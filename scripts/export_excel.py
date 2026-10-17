#!/usr/bin/env python3
"""把防水材料 JSON 导出成 Excel（汇总 / 按单体分布 / 原始清单三个工作表）。

用法:
    python3 export_excel.py input.json 输出.xlsx [--merge-by-name] [--sheet-name 新津项目]

input.json 为 extract_from_excel.py 或 extract_waterproof.py 的输出；
--merge-by-name 时按 (材料大类, 材料名称) 合并同名材料、数量求和。
"""
import argparse
import json
import os
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate_report import merge_items_by_name  # noqa: E402

HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=11)
TYPE_FILL = PatternFill('solid', fgColor='DDEBF7')
SUB_FILL = PatternFill('solid', fgColor='F2F2F2')
TOTAL_FILL = PatternFill('solid', fgColor='FFF2CC')
TITLE_FONT = Font(bold=True, size=13, color='1F4E79')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=False)

TYPE_ORDER = ['防水卷材', '防水涂料', '防水砂浆']


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(data, merge=False, sheet_name='防水材料'):
    items = data.get('items', [])
    if merge:
        items = merge_items_by_name(items)
    project = data.get('project_name') or sheet_name
    src = data.get('data_source') or ''

    wb = Workbook()

    # ---------- Sheet 1: 汇总 ----------
    ws = wb.active
    ws.title = '防水材料汇总'
    ws['A1'] = f'{project} — 防水材料汇总'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = f'数据来源：{src}' if src else ''
    ws['A2'].font = Font(size=9, color='808080')

    headers = ['序号', '材料大类', '材料名称', '单位', '工程量', '占比', '使用部位（单体）', '出现处数']
    hrow = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=hrow, column=i, value=h)
    style_header(ws, hrow, len(headers))

    grouped = OrderedDict((t, []) for t in TYPE_ORDER)
    for it in items:
        grouped.setdefault(it.get('type', '其他'), []).append(it)

    grand = sum(float(it.get('quantity') or 0) for it in items)
    r = hrow + 1
    for t, rows in grouped.items():
        if not rows:
            continue
        sub = sum(float(it.get('quantity') or 0) for it in rows)
        ws.cell(row=r, column=2, value=t)
        ws.cell(row=r, column=3, value=f'{t}小计（{len(rows)}种）')
        ws.cell(row=r, column=5, value=round(sub, 2))
        ws.cell(row=r, column=6, value=(sub / grand) if grand else 0)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = TYPE_FILL
            cell.font = Font(bold=True)
            cell.border = BORDER
            cell.alignment = LEFT if c in (2, 3) else CENTER
        ws.cell(row=r, column=6).number_format = '0.00%'
        ws.cell(row=r, column=5).number_format = '#,##0.00'
        r += 1

        for n, it in enumerate(rows, start=1):
            qty = float(it.get('quantity') or 0)
            ws.cell(row=r, column=1, value=n)
            ws.cell(row=r, column=2, value=t)
            ws.cell(row=r, column=3, value=it.get('name', ''))
            ws.cell(row=r, column=4, value=it.get('unit', 'm²'))
            ws.cell(row=r, column=5, value=round(qty, 2)).number_format = '#,##0.00'
            ws.cell(row=r, column=6, value=(qty / grand) if grand else 0).number_format = '0.00%'
            ws.cell(row=r, column=7, value=it.get('location', ''))
            ws.cell(row=r, column=8, value=it.get('_merged_n') or 1)
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                cell.alignment = LEFT if c in (2, 3, 7) else CENTER
            r += 1

    ws.cell(row=r, column=2, value='总计')
    ws.cell(row=r, column=3, value=f'合计 {len(items)} 种材料')
    ws.cell(row=r, column=5, value=round(grand, 2)).number_format = '#,##0.00'
    ws.cell(row=r, column=6, value=1).number_format = '0.00%'
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = TOTAL_FILL
        cell.font = Font(bold=True)
        cell.border = BORDER
        cell.alignment = LEFT if c in (2, 3) else CENTER

    set_widths(ws, [6, 14, 46, 6, 13, 9, 40, 10])
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    ws.auto_filter.ref = f'A{hrow}:H{r - 1}'

    # ---------- Sheet 2: 按单体分布 ----------
    ws2 = wb.create_sheet('按单体分布')
    ws2['A1'] = f'{project} — 各单体防水材料分布'
    ws2['A1'].font = TITLE_FONT

    raw_items = data.get('items', [])

    def building_of(it):
        return (str(it.get('building') or '').strip()
                or str(it.get('location') or '').split('、')[0].strip())

    buildings = []
    for it in raw_items:
        b = building_of(it)
        if b and b not in buildings:
            buildings.append(b)
    buildings.sort()

    pivot = OrderedDict()
    for it in raw_items:
        key = (it.get('type', ''), it.get('name', ''))
        row = pivot.setdefault(key, {b: 0.0 for b in buildings})
        b = building_of(it)
        if b not in row:
            row[b] = 0.0
        row[b] += float(it.get('quantity') or 0)

    heads = ['材料大类', '材料名称'] + buildings + ['合计']
    hrow2 = 3
    for i, h in enumerate(heads, start=1):
        ws2.cell(row=hrow2, column=i, value=h)
    style_header(ws2, hrow2, len(heads))

    r2 = hrow2 + 1
    col_totals = {b: 0.0 for b in buildings}
    for (t, name), row in pivot.items():
        ws2.cell(row=r2, column=1, value=t)
        ws2.cell(row=r2, column=2, value=name)
        tot = 0.0
        for i, b in enumerate(buildings, start=3):
            v = row.get(b, 0.0)
            tot += v
            col_totals[b] += v
            ws2.cell(row=r2, column=i, value=round(v, 2) if v else None)
            ws2.cell(row=r2, column=i).number_format = '#,##0.00'
        ws2.cell(row=r2, column=len(heads), value=round(tot, 2)).number_format = '#,##0.00'
        for c in range(1, len(heads) + 1):
            cell = ws2.cell(row=r2, column=c)
            cell.border = BORDER
            cell.alignment = LEFT if c in (1, 2) else CENTER
            if c == len(heads):
                cell.font = Font(bold=True)
        r2 += 1

    ws2.cell(row=r2, column=2, value='合计')
    grand2 = 0.0
    for i, b in enumerate(buildings, start=3):
        ws2.cell(row=r2, column=i, value=round(col_totals[b], 2)).number_format = '#,##0.00'
        grand2 += col_totals[b]
    ws2.cell(row=r2, column=len(heads), value=round(grand2, 2)).number_format = '#,##0.00'
    for c in range(1, len(heads) + 1):
        cell = ws2.cell(row=r2, column=c)
        cell.fill = TOTAL_FILL
        cell.font = Font(bold=True)
        cell.border = BORDER
        cell.alignment = LEFT if c == 2 else CENTER

    set_widths(ws2, [14, 46] + [12] * len(buildings) + [13])
    ws2.freeze_panes = ws2.cell(row=hrow2 + 1, column=3)

    # ---------- Sheet 3: 原始清单 ----------
    ws3 = wb.create_sheet('原始清单')
    heads3 = ['序号', '材料大类', '项目编码', '材料名称', '单体', '部位', '单位', '工程量']
    hrow3 = 1
    for i, h in enumerate(heads3, start=1):
        ws3.cell(row=hrow3, column=i, value=h)
    style_header(ws3, hrow3, len(heads3))

    raw = data.get('items', [])
    r3 = hrow3 + 1
    for n, it in enumerate(raw, start=1):
        ws3.cell(row=r3, column=1, value=n)
        ws3.cell(row=r3, column=2, value=it.get('type', ''))
        ws3.cell(row=r3, column=3, value=it.get('code', ''))
        ws3.cell(row=r3, column=4, value=it.get('name', ''))
        ws3.cell(row=r3, column=5, value=it.get('building', ''))
        ws3.cell(row=r3, column=6, value=it.get('location', ''))
        ws3.cell(row=r3, column=7, value=it.get('unit', 'm²'))
        ws3.cell(row=r3, column=8, value=round(float(it.get('quantity') or 0), 2)).number_format = '#,##0.00'
        for c in range(1, len(heads3) + 1):
            cell = ws3.cell(row=r3, column=c)
            cell.border = BORDER
            cell.alignment = LEFT if c in (2, 4, 5, 6) else CENTER
        r3 += 1

    ws3.cell(row=r3, column=4, value=f'合计 {len(raw)} 条')
    ws3.cell(row=r3, column=8, value=round(sum(float(i.get('quantity') or 0) for i in raw), 2)).number_format = '#,##0.00'
    for c in range(1, len(heads3) + 1):
        cell = ws3.cell(row=r3, column=c)
        cell.fill = TOTAL_FILL
        cell.font = Font(bold=True)
        cell.border = BORDER

    set_widths(ws3, [6, 12, 16, 46, 10, 18, 6, 13])
    ws3.freeze_panes = ws3.cell(row=hrow3 + 1, column=1)
    ws3.auto_filter.ref = f'A{hrow3}:H{r3 - 1}'

    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input')
    ap.add_argument('output')
    ap.add_argument('--merge-by-name', action='store_true')
    ap.add_argument('--sheet-name', default='防水材料')
    args = ap.parse_args()

    with open(args.input, encoding='utf-8') as f:
        data = json.load(f)
    wb = build_workbook(data, merge=args.merge_by_name, sheet_name=args.sheet_name)
    wb.save(args.output)
    print(f'Excel 已生成: {args.output}')


if __name__ == '__main__':
    main()
